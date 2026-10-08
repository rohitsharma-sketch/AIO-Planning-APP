"""
MRP Re-apportionment Engine — FastAPI Router v2
================================================
Redistributes historical sales from discontinued MRP slabs to valid (listed) MRP slabs.

Logic (the planners' MRP master, user 2026-10-08 - "MRP MASTER - 27-04-2026 - New(Shubham).xlsx", sheet "MRP Adj"):
  - The master lists, per Department x Display, every old MRP (MRP_CURRENT) and its new MRP (MRP_LISTED; 0 =
    discontinued).
  - Sales at a listed MRP move to its MRP_LISTED.
  - Sales at a discontinued MRP split between the nearest listed MRP below and the nearest above (old MRPs in value
    order, same Department x Display): Summer 60% down / 40% up, every other attribute 40% down / 60% up (Regular and
    Occasional per the master's assumptions; PreWinter / Winter the same, user 2026-10-08). Only one side listed ->
    100% there. Nothing listed in the Department x Display -> Unmapped.
  - Totals are preserved exactly per Store x Dept x Display x Attribute x Month (Unmapped rows reported apart).

Sources:
  Sales:       the suite's sales engine (user, 2026-09-29: "LY actual sales will be taken from the sales engine
               which is embedded in all the other apps, only the new MRP structure will be given") - the Calendar
               engine's month-wise data-lake file, read with its own reader, LY Mar-Jun 2026, and checked cell by
               cell against the Calendar department snapshot (store x department x month) Sales Plan reads.
  MRP Mapping: MRP Merging Engine/Sales Reapportionment/MRP Mapping/MRP Mapping Master.xlsx (the only upload)

Endpoints:
  GET  /status           → file presence, last-run metadata
  GET  /sales-status     → sales engine file, LY months, tie to Calendar dept sales
  GET  /checks           → pre-run checks (red = Run disabled), each with how to fix it
  GET  /mrp-groups       → groups (Dept, Display): listed MRPs and where each discontinued one goes
  POST /run              → redistribution by the nearest-listed-MRP rule
  GET  /download         → latest output Excel
  POST /mapping/upload   → new version of the MRP master (checked first; the old one moves to MRP Mapping/Archive)
  GET  /mapping/download → the active MRP master, as uploaded
"""

import os
import io
import glob
import json
import shutil
import tempfile
import warnings
import datetime
import threading
from typing import Optional, Dict, List

import sys
import pandas as pd
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from apportion import SHOWN  # noqa: E402
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from plan_cache import save_json  # noqa: E402 - atomic JSON writes (audit 2026-10-06)

warnings.filterwarnings("ignore")

router = APIRouter()

# ── Paths ──────────────────────────────────────────────────────────────────────
BASE         = r"C:\Users\A9820\Documents\CLaude - New Projects"
MAPPING_DIR  = os.path.join(BASE, r"MRP Merging Engine\Sales Reapportionment\MRP Mapping")
OUTPUT_DIR   = os.path.join(BASE, r"MRP Merging Engine\Sales Reapportionment\Output")
LAST_RUN_JSON = os.path.join(OUTPUT_DIR, "_last_run.json")

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(MAPPING_DIR, exist_ok=True)


# ── Pydantic models ────────────────────────────────────────────────────────────

class RunRequest(BaseModel):
    pass   # the split is fixed by the master's rule (typed %s per group until 2026-10-08)


# ── Helpers ────────────────────────────────────────────────────────────────────

def _norm_cols(df: pd.DataFrame) -> pd.DataFrame:
    df.columns = [str(c).strip().upper().replace(" ", "_").replace("-", "_") for c in df.columns]
    return df


# LY months = the plan's LY (Mar-Jun 2026, as BIS / AOP); labels as the old sales template ("Mar 2026").
LY_MONTHS = ["2026-03", "2026-04", "2026-05", "2026-06"]
_SALES_CACHE: dict = {}   # {"key": (path, mtime), "df", "month_cols", "check"}
_SALES_LOCK = threading.Lock()   # one read at a time: a page opened during the startup warm-up waits for it


def _calendar_scans():
    """The Calendar engine's own data-lake reader (RS Planning Platform/backend/calendar_engine/scans.py)."""
    import sys
    plat = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "RS Planning Platform", "backend")
    if plat not in sys.path:
        sys.path.insert(0, plat)
    from calendar_engine import scans
    return scans


def _find_sales_file() -> Optional[str]:
    files = _calendar_scans()._latest_monthwise_files()
    return files[0] if files else None


def _engine_check(raw: pd.DataFrame) -> dict:
    """Store x department x month of the MRP-level sales vs the Calendar 'actual_dept' snapshot (Rs in, lakh diff)."""
    import collections
    import sys
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "Tentative AOP Forecaster"))
    from db.base import SessionLocal
    from sqlalchemy import text
    with SessionLocal() as db:
        snap = db.execute(text("SELECT rows FROM calendar.sales_snapshots "
                               "WHERE source_type='mw' AND kind='actual_dept'")).scalar() or []
    want = set(LY_MONTHS)
    cal = collections.defaultdict(float)
    for r in snap:
        if r["col"] in want:
            cal[(str(r["store"]).strip().upper(), str(r["DEPARTMENT"]).strip().upper(), r["col"])] += float(r["value"])
    ours = raw.groupby(["STORE", "DEPARTMENT", "ym"])["SL_V"].sum().to_dict()
    keys = set(cal) | set(ours)
    worst = max((abs(ours.get(k, 0.0) - cal.get(k, 0.0)) / 1e5 for k in keys), default=0.0)
    return {"cells": len(keys), "max_diff_lakh": round(worst, 6), "pass": bool(keys) and worst <= 0.01,
            "total_lakh": round(sum(ours.values()) / 1e5, 2), "calendar_total_lakh": round(sum(cal.values()) / 1e5, 2)}


def _engine_sales() -> tuple:
    """LY sales at Store x Division x Dept x MRP x Display x Attribute, one column per LY month, from the sales
    engine's file. ponytail: one cached copy, re-read only when the data-lake file changes."""
    scans = _calendar_scans()
    path = _find_sales_file()
    if not path:
        raise ValueError(f"Sales engine file not reachable: {scans.PARQUET_DIR}")
    key = (path, os.path.getmtime(path))
    with _SALES_LOCK:
        if _SALES_CACHE.get("key") != key:
            _read_engine_sales(scans, path, key)
    return _SALES_CACHE["df"].copy(), list(_SALES_CACHE["month_cols"]), _SALES_CACHE["check"]


def _read_engine_sales(scans, path, key):
    lo, hi = scans._month_bounds(LY_MONTHS)
    raw = scans._read_file_filtered(path, ["BILLMONTH", "DIVISION", "STORE_NAME", "SL_V", "DEPARTMENT", "MRP",
                                           "DISPLAY_TYPE", "ATTRIBUTE1"], "BILLMONTH", lo, hi)
    raw["ym"] = raw["BILLMONTH"].dt.strftime("%Y-%m")
    raw = raw[raw["ym"].isin(LY_MONTHS)].rename(
        columns={"STORE_NAME": "STORE", "DISPLAY_TYPE": "DISPLAY", "ATTRIBUTE1": "ATTRIBUTE"})
    for c in ("STORE", "DIVISION", "DEPARTMENT", "DISPLAY", "ATTRIBUTE"):
        raw[c] = raw[c].astype(str).str.strip().str.upper()
    raw["MRP"] = pd.to_numeric(raw["MRP"], errors="coerce").fillna(0).round().astype(int)
    raw["SL_V"] = pd.to_numeric(raw["SL_V"], errors="coerce").fillna(0.0)
    check = _engine_check(raw)
    label = {m: datetime.date(int(m[:4]), int(m[5:]), 1).strftime("%b %Y") for m in LY_MONTHS}
    wide = raw.pivot_table(index=["STORE", "DIVISION", "DEPARTMENT", "MRP", "DISPLAY", "ATTRIBUTE"],
                           columns="ym", values="SL_V", aggfunc="sum", fill_value=0.0).reset_index()
    wide.columns.name = None
    wide = wide.rename(columns=label)
    month_cols = [label[m] for m in LY_MONTHS if label[m] in wide.columns]
    _SALES_CACHE.update(key=key, df=wide, month_cols=month_cols, check=check)


def warm_sales():
    """Read the sales engine file once at server start (user, 2026-10-07: the first MRP Re-apportionment open after a
    restart took ~85 s). Never raises - the page shows the reason itself if the file isn't reachable."""
    try:
        _engine_sales()
    except Exception as e:  # noqa: BLE001
        print(f"MRP re-apportionment warm-up skipped: {e}", flush=True)


def _mapped_sales(mapping_df: pd.DataFrame) -> tuple:
    """Engine sales of the departments the MRP structure covers (the rest of the business is not re-apportioned)."""
    df, month_cols, check = _engine_sales()
    return df[df["DEPARTMENT"].isin(set(mapping_df["DEPARTMENT"]))].reset_index(drop=True), month_cols, check


def _find_mapping_file() -> Optional[str]:
    files = []
    for pat in ("*.xlsx", "*.xls"):
        files.extend(glob.glob(os.path.join(MAPPING_DIR, pat)))
    return max(files, key=os.path.getmtime) if files else None


def _wide_master(path: str) -> Optional[pd.DataFrame]:
    """The master as the planners keep it: one column per Department x Display x old MRP, rows labelled DEPARTMENT /
    DISPLAY / MRP_CURRENT / MRP_LISTED in column A (sheet "MRP Adj"). None if no sheet is laid out that way."""
    need = ("DEPARTMENT", "DISPLAY", "MRP_CURRENT", "MRP_LISTED")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = {str(r[0]).strip().upper(): r for r in ws.iter_rows(max_row=30, values_only=True)
                    if r and r[0] is not None}
            if all(k in rows for k in need):
                n = min(len(rows[k]) for k in need)
                df = pd.DataFrame({k: list(rows[k][1:n]) for k in need})
                return df[df["DEPARTMENT"].notna() & pd.to_numeric(df["MRP_CURRENT"], errors="coerce").notna()]
    finally:
        wb.close()
    return None


def _load_mapping(path: Optional[str] = None) -> pd.DataFrame:
    """Load the MRP master -> DEPARTMENT, DISPLAY, MRP_CURRENT, MRP_LISTED (one row per Department x Display x old
    MRP). Reads the planners' wide "MRP Adj" layout, or a plain list with those columns."""
    path = path or _find_mapping_file()
    if not path:
        raise ValueError(f"No mapping file found in: {MAPPING_DIR}")

    df = _wide_master(path)
    if df is None:
        df = _norm_cols(pd.read_excel(path))
        rename = {}
        for c in df.columns:
            cu = c.upper()
            if cu in ("DEPARTMENT", "DEPT"):               rename[c] = "DEPARTMENT"
            elif cu in ("DISPLAY", "DISPLAY_TYPE"):        rename[c] = "DISPLAY"
            elif cu in ("MRP_CURRENT", "CURRENT_MRP",
                        "MRP", "RAW_MRP", "OLD_MRP"):      rename[c] = "MRP_CURRENT"
            elif cu in ("MRP_LISTED", "LISTED_MRP",
                        "NEW_MRP", "FINAL_MRP"):            rename[c] = "MRP_LISTED"
        df.rename(columns=rename, inplace=True)
        missing = [c for c in ("DEPARTMENT", "DISPLAY", "MRP_CURRENT", "MRP_LISTED") if c not in df.columns]
        if missing:
            raise ValueError(f"MRP Mapping Master missing columns: {missing}")

    df = df.dropna(subset=["DEPARTMENT"])[["DEPARTMENT", "DISPLAY", "MRP_CURRENT", "MRP_LISTED"]].copy()
    for c in ["DEPARTMENT", "DISPLAY"]:
        df[c] = df[c].astype(str).str.strip().str.upper()
    df["MRP_CURRENT"] = pd.to_numeric(df["MRP_CURRENT"], errors="coerce").fillna(0).round().astype(int)
    df["MRP_LISTED"]  = pd.to_numeric(df["MRP_LISTED"],  errors="coerce").fillna(0).round().astype(int)
    df = df.drop_duplicates()
    clash = df[df.duplicated(["DEPARTMENT", "DISPLAY", "MRP_CURRENT"], keep=False)]
    if not clash.empty:
        raise ValueError(f"MRP master gives {len(clash)} old MRPs two different new MRPs, "
                         f"e.g. {clash.head(3).to_dict('records')}")
    return df.reset_index(drop=True)


# Share of a discontinued MRP's sales that goes to the nearest listed MRP BELOW it; the rest goes to the nearest above.
LOWER_PCT = {"SUMMER": 60.0}
LOWER_PCT_DEFAULT = 40.0   # Regular / Occasional (master's assumptions); PreWinter / Winter too (user, 2026-10-08)
KEYS = ["DEPARTMENT", "DISPLAY", "MRP_CURRENT"]


def _targets(mapping_df: pd.DataFrame) -> pd.DataFrame:
    """Each master row plus BELOW / ABOVE: the new MRP of the nearest listed old MRP below / above it in the same
    Department x Display (0 = none on that side) - what the master's Initial1 / Initial2 rows work out."""
    out = []
    for _, g in mapping_df.sort_values("MRP_CURRENT").groupby(["DEPARTMENT", "DISPLAY"], sort=False):
        lst = g["MRP_LISTED"].tolist()
        below, run = [], 0
        for v in lst:
            below.append(run)
            run = v or run
        above, run = [], 0
        for v in reversed(lst):
            above.append(run)
            run = v or run
        out.append(g.assign(BELOW=below, ABOVE=above[::-1]))
    cols = ["DEPARTMENT", "DISPLAY", "MRP_CURRENT", "MRP_LISTED", "BELOW", "ABOVE"]
    return pd.concat(out, ignore_index=True)[cols] if out else pd.DataFrame(columns=cols)


def _redistribute(sales_df: pd.DataFrame, mapping_df: pd.DataFrame, month_cols: list) -> tuple:
    """Listed MRPs move to their new MRP; a discontinued MRP splits between the nearest listed MRPs below / above
    (LOWER_PCT by attribute, 100% to the only side). Returns (output_df, unmapped_df, log_lines)."""
    s = sales_df[(sales_df[month_cols] != 0).any(axis=1)].rename(columns={"MRP": "MRP_CURRENT"})
    s = s.assign(MRP_CURRENT=s["MRP_CURRENT"].astype(int)).merge(_targets(mapping_df), on=KEYS, how="left")
    not_in = s["MRP_LISTED"].isna()
    listed = ~not_in & (s["MRP_LISTED"] > 0)
    nowhere = ~not_in & ~listed & (s["BELOW"] == 0) & (s["ABOVE"] == 0)
    d = s[~not_in & ~listed & ~nowhere]

    lo_pct = d["ATTRIBUTE"].map(LOWER_PCT).fillna(LOWER_PCT_DEFAULT)
    lo_pct = lo_pct.where(d["ABOVE"] > 0, 100.0).where(d["BELOW"] > 0, 0.0)
    lo = d.assign(LISTED_MRP=d["BELOW"], SHARE_PCT=lo_pct)
    lo[month_cols] = d[month_cols].mul(lo_pct / 100.0, axis=0)
    hi = d.assign(LISTED_MRP=d["ABOVE"], SHARE_PCT=100.0 - lo_pct)
    hi[month_cols] = d[month_cols] - lo[month_cols]   # the rest, so the two parts add back to the month exactly

    key_cols = ["STORE", "DIVISION", "DEPARTMENT", "DISPLAY", "ATTRIBUTE", "MRP_CURRENT", "LISTED_MRP"]
    parts = [s[listed].assign(LISTED_MRP=s["MRP_LISTED"], SHARE_PCT=100.0), lo[lo_pct > 0], hi[lo_pct < 100]]
    output_df = pd.concat(parts, ignore_index=True)[key_cols + ["SHARE_PCT"] + month_cols]
    output_df["LISTED_MRP"] = output_df["LISTED_MRP"].astype(int)
    # one row per store x ... x old MRP x new MRP (both neighbours can carry the same new MRP: their shares add)
    output_df = output_df.groupby(key_cols, dropna=False)[["SHARE_PCT"] + month_cols].sum().reset_index()

    unmapped_df = pd.concat([s[not_in].assign(REASON="MRP combination not in mapping master"),
                             s[nowhere].assign(REASON="No listed MRP in this Department x Display")],
                            ignore_index=True)
    unmapped_df = unmapped_df.assign(LISTED_MRP=0)[key_cols[:-1] + ["LISTED_MRP"] + month_cols + ["REASON"]]

    log = [f"Rule: discontinued MRP -> nearest listed MRP below / above; Summer {LOWER_PCT['SUMMER']:g}/"
           f"{100 - LOWER_PCT['SUMMER']:g}, others {LOWER_PCT_DEFAULT:g}/{100 - LOWER_PCT_DEFAULT:g}; one side -> 100%",
           f"Listed rows: {int(listed.sum()):,} | split rows: {len(d):,} | unmapped rows: {len(unmapped_df):,}"]
    for (dept, disp), n in unmapped_df.groupby(["DEPARTMENT", "DISPLAY"]).size().items():
        log.append(f"WARN: {n} rows unmapped in [{dept} | {disp}]")
    log.append(f"Output rows: {len(output_df)}")
    return output_df, unmapped_df, log


def _validate(sales_df: pd.DataFrame, output_df: pd.DataFrame,
              month_cols: list, unmapped_df: Optional[pd.DataFrame] = None) -> tuple[list[dict], bool]:
    """Store x Department totals before vs after. Unmapped rows count on the 'after' side - they are reported on their
    own sheet, not lost (user, 2026-10-08: departments with no listed MRP go to Unmapped)."""
    if unmapped_df is not None and not unmapped_df.empty:
        output_df = pd.concat([output_df, unmapped_df[["STORE", "DEPARTMENT"] + month_cols]], ignore_index=True)
    tol = SHOWN   # any difference that shows at 8 decimals (was 0.01)
    val_rows = []
    val_pass = True

    # Sum all month sales per row for Store×Dept totals
    sales_df = sales_df.copy()
    output_df = output_df.copy() if not output_df.empty else pd.DataFrame()

    sales_df["TOTAL"] = sales_df[month_cols].sum(axis=1)
    orig = sales_df.groupby(["STORE", "DEPARTMENT"])["TOTAL"].sum().reset_index()
    orig.rename(columns={"TOTAL": "BEFORE"}, inplace=True)

    if output_df.empty:
        for _, r in orig.iterrows():
            val_rows.append({"STORE": r["STORE"], "DEPARTMENT": r["DEPARTMENT"],
                             "BEFORE": r["BEFORE"], "AFTER": 0.0,
                             "DIFF": r["BEFORE"], "STATUS": "FAIL"})
        return val_rows, False

    mc_in_output = [c for c in month_cols if c in output_df.columns]
    output_df["TOTAL"] = output_df[mc_in_output].sum(axis=1)
    after = output_df.groupby(["STORE", "DEPARTMENT"])["TOTAL"].sum().reset_index()
    after.rename(columns={"TOTAL": "AFTER"}, inplace=True)

    merged = orig.merge(after, on=["STORE", "DEPARTMENT"], how="outer").fillna(0)
    merged["DIFF"]   = (merged["AFTER"] - merged["BEFORE"]).abs()
    merged["STATUS"] = merged["DIFF"].apply(lambda d: "PASS" if d <= tol else "FAIL")
    val_pass = bool((merged["STATUS"] == "PASS").all())
    return merged.to_dict("records"), val_pass


# ── Excel output builder ───────────────────────────────────────────────────────

def _build_excel(output_df, unmapped_df, val_rows, val_pass, log_lines, sales_df, month_cols) -> bytes:
    """The run's workbook. xlsxwriter with one format per column and the row banding as a single conditional format
    (user, 2026-10-08: "speed up the excel export" - openpyxl styled ~3.4 M cells one by one, ~7 min a run)."""
    import xlsxwriter

    NAVY = "#1F3864"; GREEN = "#C6EFCE"; RED = "#FF9999"; GREY = "#2C3538"; DKGREY = "#1C2022"
    INT_COLS = {"MRP_CURRENT", "LISTED_MRP", "SHARE_PCT"}
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False})
    base = {"font_name": "Arial", "font_size": 9, "border": 1, "border_color": "#2D3B40", "valign": "vcenter"}

    def F(**k):
        return wb.add_format({**base, **k})

    head = F(bold=True, font_color="#FFFFFF", font_size=10, bg_color=NAVY, align="center", text_wrap=True)
    band = wb.add_format({"bg_color": GREY})
    col_fmt = {kind: F(bg_color=DKGREY, font_color="#E2EAED", align="left" if kind is None else "right",
                       num_format={None: "General", "int": "#,##0", "dec": "#,##0.00000000"}[kind])
               for kind in (None, "int", "dec")}
    note = wb.add_format({"font_name": "Arial", "font_size": 9, "font_color": "#7FA0AD"})

    def write_df(ws, df, title):
        if df.empty:
            ws.write(0, 0, f"No data — {title}", note)
            return
        cols = list(df.columns)
        ws.set_row(0, 26)
        ws.write_row(0, 0, cols, head)
        data = []
        for ci, c in enumerate(cols):
            num = pd.api.types.is_numeric_dtype(df[c])
            kind = ("int" if c in INT_COLS else "dec") if num else None
            ws.set_column(ci, ci, min(28, max(10, len(c) + 4)), col_fmt[kind])
            v = df[c].astype(float).round(8) if num else df[c]   # full precision to 8 dp, as before
            data.append(v.astype(object).where(v.notna(), None).tolist())
        for r, row in enumerate(zip(*data), 1):
            ws.write_row(r, 0, row)
        last = len(df)
        ws.conditional_format(1, 0, last, len(cols) - 1, {"type": "formula", "criteria": "=MOD(ROW(),2)=0", "format": band})
        ws.freeze_panes(1, 0)
        ws.autofilter(0, 0, last, len(cols) - 1)

    write_df(wb.add_worksheet("Reapportioned Sales"), output_df, "Reapportioned Sales")

    ws2 = wb.add_worksheet("Validation")
    ws2.write(0, 0, "Validation — Store × Department Totals", F(bold=True, font_color="#E2EAED", font_size=12, border=0))
    ws2.write(1, 0, f"Overall: {'PASSED' if val_pass else 'FAILED'}",
              F(bold=True, font_color="#006100" if val_pass else "#9C0006", font_size=11, border=0))
    vcols = ["STORE", "DEPARTMENT", "BEFORE", "AFTER", "DIFF", "STATUS"]
    ws2.set_row(3, 26)
    ws2.write_row(3, 0, vcols, head)
    vf = {(ok, ci): F(bg_color=GREEN if ok else RED, font_color="#000000", bold=not ok,
                      align="right" if ci >= 2 else "left",
                      num_format="#,##0.00000000" if ci in (2, 3, 4) else "General")
          for ok in (True, False) for ci in range(6)}
    for ri, r in enumerate(val_rows, 4):
        ok = r["STATUS"] == "PASS"
        for ci, k in enumerate(vcols):
            ws2.write(ri, ci, r[k], vf[(ok, ci)])

    write_df(wb.add_worksheet("Unmapped (Listed MRP 0)"), unmapped_df, "Unmapped Sales")

    ws4 = wb.add_worksheet("Engine Log")
    ws4.set_column(0, 0, 90)
    lf = {c: wb.add_format({"font_name": "Arial", "font_size": 9, "font_color": c})
          for c in ("#7FA0AD", "#9C0006", "#006100", "#9C6500", "#E2EAED")}
    ws4.write(0, 0, f"Run: {datetime.datetime.now().strftime('%d-%b-%Y %H:%M:%S')}", lf["#7FA0AD"])
    for ri, line in enumerate(log_lines, 2):
        c = ("#9C0006" if "FAIL" in line or "ERROR" in line else "#006100" if "PASS" in line
             else "#9C6500" if "WARN" in line else "#E2EAED")
        ws4.write_string(ri, 0, line, lf[c])

    write_df(wb.add_worksheet("Original Sales"), sales_df, "Original Sales")

    wb.close()
    return buf.getvalue()


# ── Endpoints ──────────────────────────────────────────────────────────────────

# ── Checks (user, 2026-10-08: "how can checks be embedded in this - and how can we rerun if the checks are not passed ?")
# Each check: ok / warn / fail, what it found, and - when not ok - what to fix before pressing Run again. A failed
# pre-run check stops the run (the endpoints refuse it too); a failed post-run check marks the run FAILED.

FIX_ENGINE = "An engine fault, not your data - this run is marked FAILED; report it, and Run again once it is fixed."


def _check(key, label, status, detail, fix=""):
    return {"key": key, "label": label, "status": status, "detail": detail, "fix": fix if status != "ok" else ""}


def _pre_checks() -> list:
    out, mp, sales = [], None, None
    try:
        mp = _load_mapping()
        t = _targets(mp)
        top = t.groupby(["DEPARTMENT", "DISPLAY"])["MRP_LISTED"].max()
        out.append(_check("master", "MRP master", "ok",
                          f"{os.path.basename(_find_mapping_file())}: {len(mp):,} old MRPs in {len(top)} Dept x Display"))
    except Exception as e:
        out.append(_check("master", "MRP master", "fail", str(e),
                          "Import a usable MRP master (MRP Mapping Master card), then Run again."))
    try:
        sales, months, check = _mapped_sales(mp) if mp is not None else _engine_sales()
        ok = bool(check["pass"])
        out.append(_check("sales_tie", "LY sales tie to the Calendar department sales", "ok" if ok else "fail",
                          f"{check['cells']:,} store x dept x month cells, max diff {check['max_diff_lakh']} L",
                          "The sales engine file and the Calendar sales differ - wait for the nightly sales sync (or "
                          "Sync now on Landing), press Refresh, then Run again."))
    except Exception as e:
        out.append(_check("sales_tie", "LY sales tie to the Calendar department sales", "fail", str(e),
                          "The sales engine file isn't reachable - check the data-lake share, press Refresh, then Run again."))
        sales = None
    if mp is not None and sales is not None:
        x = sales.rename(columns={"MRP": "MRP_CURRENT"}).merge(t, on=KEYS, how="left")
        v = x[months].sum(axis=1)
        total, miss = float(v.sum()), float(v[x["MRP_LISTED"].isna()].sum())
        nowhere = float(v[(x["MRP_LISTED"] == 0) & (x["BELOW"] == 0) & (x["ABOVE"] == 0)].sum())
        pct = 100 * (total - miss) / total if total else 100.0
        out.append(_check("coverage", "LY sales on MRPs in the master", "ok" if miss == 0 else "warn",
                          f"{(total - miss) / 1e5:,.2f} L of {total / 1e5:,.2f} L ({pct:.2f}%) of the master's departments; "
                          f"{miss / 1e5:,.2f} L on old MRPs not in the master go to Unmapped",
                          "Add the missing old MRPs to the master (Export current, edit, Import new version), then Run again."))
        n = int(top.eq(0).sum())
        out.append(_check("nothing_listed", "Dept x Display with nothing listed", "ok" if n == 0 else "warn",
                          f"{n} of {len(top)}; {nowhere / 1e5:,.2f} L of LY sales go to Unmapped",
                          "List at least one MRP in those groups in the master, Import it, then Run again."))
    return out


def _post_checks(sales_df, output_df, unmapped_df, month_cols, mapping_df, val_pass) -> list:
    out = [_check("store_dept", "Store x Dept totals: before = after + Unmapped", "ok" if val_pass else "fail",
                  "every store x department to 8 decimals" if val_pass else "some store x department totals differ "
                  "(see the Validation sheet)", FIX_ENGINE)]
    after = output_df[month_cols].sum() + (unmapped_df[month_cols].sum() if not unmapped_df.empty else 0)
    # in lakh, as the suite's other ties: a ~Rs 1,000 Cr month summed over 140k rows carries ~Rs 1e-7 of float noise
    worst = float((sales_df[month_cols].sum() - after).abs().max()) / 1e5
    out.append(_check("months", "Month totals: before = after + Unmapped", "ok" if worst <= SHOWN else "fail",
                      f"{len(month_cols)} months, largest difference {worst:.8f} L", FIX_ENGINE))
    k = ["STORE", "DIVISION", "DEPARTMENT", "DISPLAY", "ATTRIBUTE", "MRP_CURRENT"]
    shares = output_df.groupby(k, dropna=False)["SHARE_PCT"].sum()
    bad = int(((shares - 100).abs() > 1e-9).sum())
    out.append(_check("shares", "Each old MRP's shares add to 100%", "ok" if bad == 0 else "fail",
                      f"{len(shares):,} store x old MRP rows" + (f", {bad:,} not at 100%" if bad else ""), FIX_ENGINE))
    listed = mapping_df[mapping_df["MRP_LISTED"] > 0][["DEPARTMENT", "DISPLAY", "MRP_LISTED"]].drop_duplicates()
    hit = output_df[["DEPARTMENT", "DISPLAY", "LISTED_MRP"]].drop_duplicates().merge(
        listed.rename(columns={"MRP_LISTED": "LISTED_MRP"}), how="left", indicator=True)
    stray = int((hit["_merge"] == "left_only").sum())
    out.append(_check("targets", "Every new MRP is a listed MRP of its Dept x Display", "ok" if stray == 0 else "fail",
                      f"{len(hit):,} Dept x Display x new MRP" + (f", {stray} not listed in the master" if stray else ""),
                      FIX_ENGINE))
    return out


@router.get("/checks")
def get_checks():
    """The pre-run checks the page shows above Run (a red one disables it)."""
    return {"checks": _pre_checks()}


@router.get("/status")
def get_status():
    sales_file   = _find_sales_file()
    mapping_file = _find_mapping_file()
    output_files = sorted(
        glob.glob(os.path.join(OUTPUT_DIR, "MRP Reapportioned*.xlsx")),
        key=os.path.getmtime, reverse=True
    )
    last_run = None
    if os.path.exists(LAST_RUN_JSON):
        try:
            with open(LAST_RUN_JSON) as f:
                last_run = json.load(f)
        except Exception:
            pass

    return {
        "sales_file_found":   sales_file is not None,
        "sales_file":         os.path.basename(sales_file) if sales_file else None,
        "mapping_file_found": mapping_file is not None,
        "mapping_file":       os.path.basename(mapping_file) if mapping_file else None,
        "output_found":       len(output_files) > 0,
        "last_output":        os.path.basename(output_files[0]) if output_files else None,
        "last_run":           last_run,
        "mapping_folder":     MAPPING_DIR,
    }


@router.get("/sales-status")
def get_sales_status():
    """The sales engine's file, the LY months read and the tie to the Calendar department snapshot."""
    try:
        f = _find_sales_file()
        _, month_cols, check = _engine_sales()
    except Exception as e:
        return {"file_found": False, "source": "Sales engine (data lake)", "error": str(e)}
    return {
        "file_found": True,
        "source":     "Sales engine (data lake)",
        "filename":   os.path.basename(f),
        "modified":   datetime.datetime.fromtimestamp(os.path.getmtime(f)).strftime("%d %b %Y %H:%M"),
        "months":     month_cols,
        "check":      check,
    }


@router.get("/mrp-groups")
def get_mrp_groups():
    """Per Department x Display: the listed MRPs (old -> new) and, for each discontinued old MRP, the nearest listed
    new MRP below / above it (0 = none) plus how many sales rows sit on discontinued MRPs."""
    try:
        t = _targets(_load_mapping())
    except Exception as e:
        raise HTTPException(400, detail=str(e))

    disc_rows: dict = {}
    if _find_sales_file():
        try:
            sales_df, _, _ = _mapped_sales(t)
            hit = sales_df.rename(columns={"MRP": "MRP_CURRENT"}).merge(t[t["MRP_LISTED"] == 0][KEYS], on=KEYS)
            disc_rows = hit.groupby(["DEPARTMENT", "DISPLAY"]).size().to_dict()
        except Exception:
            pass

    result = []
    for (dept, display), g in t.groupby(["DEPARTMENT", "DISPLAY"]):
        on = g[g["MRP_LISTED"] > 0]
        off = g[g["MRP_LISTED"] == 0]
        result.append({
            "dept": dept, "display": display, "group_key": f"{dept}|{display}",
            "listed": [{"mrp": int(r.MRP_CURRENT), "listed": int(r.MRP_LISTED)} for r in on.itertuples()],
            "disc": [{"mrp": int(r.MRP_CURRENT), "below": int(r.BELOW), "above": int(r.ABOVE)} for r in off.itertuples()],
            "disc_rows_in_sales": int(disc_rows.get((dept, display), 0)),
        })
    return {"groups": result, "total_groups": len(result),
            "rule": {"lower_pct": LOWER_PCT, "lower_pct_default": LOWER_PCT_DEFAULT}}


@router.post("/run")
def run_engine(body: RunRequest):
    """Run re-apportionment by the master's nearest-listed-MRP rule (see module docstring)."""
    sales_path = _find_sales_file()
    if not sales_path:
        raise HTTPException(400, detail="Sales engine file not reachable")

    mapping_file = _find_mapping_file()
    if not mapping_file:
        raise HTTPException(400, detail=f"No mapping master found in: {MAPPING_DIR}")

    try:
        mapping_df = _load_mapping()
        sales_df, month_cols, check = _mapped_sales(mapping_df)
    except Exception as e:
        raise HTTPException(400, detail=str(e))
    if not check["pass"]:   # never re-apportion sales that don't tie to the sales engine's structure
        raise HTTPException(409, detail=f"Sales do not tie to the Calendar department sales (max diff {check['max_diff_lakh']} L)")

    try:
        output_df, unmapped_df, eng_log = _redistribute(sales_df, mapping_df, month_cols)
    except Exception as e:
        raise HTTPException(500, detail=f"Redistribution error: {e}")

    val_rows, val_pass = _validate(sales_df, output_df, month_cols, unmapped_df)
    post_checks = _post_checks(sales_df, output_df, unmapped_df, month_cols, mapping_df, val_pass)
    checks_pass = all(c["status"] == "ok" for c in post_checks)

    total_before = round(float(sales_df[month_cols].sum().sum()), 8)
    total_after  = round(float(output_df[month_cols].sum().sum()), 8) if not output_df.empty else 0.0
    total_unmapped = round(float(unmapped_df[month_cols].sum().sum()), 8) if not unmapped_df.empty else 0.0

    full_log = [
        f"Sales       : sales engine {os.path.basename(sales_path)} (tie to Calendar dept sales: max {check['max_diff_lakh']} L)",
        f"Mapping file: {os.path.basename(mapping_file)}",
        f"Month cols  : {', '.join(month_cols)}",
        f"Input rows  : {len(sales_df):,}",
        f"Output rows : {len(output_df):,}",
        f"Unmapped    : {len(unmapped_df)} rows, {total_unmapped:,.2f}",
        f"Total before: {total_before:,.2f}",
        f"Total after : {total_after:,.2f}",
        f"Diff        : {abs(total_after + total_unmapped - total_before):.6f} (after + unmapped vs before)",
        f"Validation  : {'PASSED' if val_pass else 'FAILED'}",
        f"Checks      : {'PASSED' if checks_pass else 'FAILED - not for use'}",
        *[f"  {'PASS' if c['status'] == 'ok' else 'FAIL'}  {c['label']}: {c['detail']}" for c in post_checks],
        "---",
        *eng_log,
    ]

    ts_str = datetime.datetime.now().strftime("%d%b%Y_%H%M")
    out_filename = f"MRP Reapportioned {ts_str}.xlsx"
    out_path = os.path.join(OUTPUT_DIR, out_filename)
    xlsx_bytes = _build_excel(output_df, unmapped_df, val_rows, val_pass, full_log, sales_df, month_cols)
    with open(out_path, "wb") as f:
        f.write(xlsx_bytes)

    preview = []
    if not output_df.empty:
        sort_cols = ["STORE", "DEPARTMENT", "DISPLAY", "ATTRIBUTE", "LISTED_MRP"]
        sc = [c for c in sort_cols if c in output_df.columns]
        preview_df = output_df.sort_values(sc).head(50)
        preview = preview_df.round(2).to_dict("records")

    run_meta = {
        "run_at":      datetime.datetime.now().isoformat(),
        "sales_file":  os.path.basename(sales_path),
        "sales_check": check,
        "input_rows":  len(sales_df),
        "output_rows": len(output_df),
        "unmapped":    len(unmapped_df),
        "total_before": total_before,
        "total_after":  total_after,
        "val_pass":     val_pass,
        "checks_pass":  checks_pass,
        "output_file":  out_filename,
        "stores":       int(sales_df["STORE"].nunique()),
        "departments":  int(sales_df["DEPARTMENT"].nunique()),
        "month_cols":   month_cols,
    }
    save_json(LAST_RUN_JSON, run_meta, indent=2)

    return {
        "ok":           True,
        "input_rows":   len(sales_df),
        "output_rows":  len(output_df),
        "unmapped":     len(unmapped_df),
        "total_before": total_before,
        "total_after":  total_after,
        "total_unmapped": total_unmapped,
        "diff":         round(abs(total_after + total_unmapped - total_before), 8),
        "val_pass":     val_pass,
        "checks_pass":  checks_pass,
        "post_checks":  post_checks,
        "val_rows":     val_rows,
        "output_file":  out_filename,
        "log":          full_log,
        "preview":      preview,
        "month_cols":   month_cols,
    }


@router.get("/download")
def download_latest():
    files = sorted(
        glob.glob(os.path.join(OUTPUT_DIR, "MRP Reapportioned*.xlsx")),
        key=os.path.getmtime, reverse=True
    )
    if not files:
        raise HTTPException(404, detail="No output file found. Run the engine first.")
    path = files[0]
    with open(path, "rb") as f:
        data = f.read()
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={os.path.basename(path)}"},
    )


# ── MRP master versions: import / export (user, 2026-10-08: "what if i have a new version for mrp mapping master ?
# Importing and exporting feature - add it") ─────────────────────────────────────────────────────────────────────

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _mapping_changes(old: Optional[pd.DataFrame], new: pd.DataFrame) -> dict:
    """Old MRPs added / removed / given a different new MRP, versus the master being replaced."""
    if old is None:
        return {"added": len(new), "removed": 0, "changed": 0}
    o, n = old.set_index(KEYS)["MRP_LISTED"], new.set_index(KEYS)["MRP_LISTED"]
    both = o.index.intersection(n.index)
    return {"added": len(n.index.difference(o.index)), "removed": len(o.index.difference(n.index)),
            "changed": int((o[both] != n[both]).sum())}


@router.post("/mapping/upload")
def upload_mapping(file: UploadFile = File(...)):
    """A new version of the MRP master. Read with the same reader a run uses before it goes live; the version it
    replaces moves to MRP Mapping\\Archive with a time stamp (kept, never overwritten)."""
    name = os.path.basename(file.filename or "").strip()
    if not name.lower().endswith(".xlsx"):
        raise HTTPException(400, detail="Upload the MRP master as an .xlsx file")
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    try:
        with os.fdopen(fd, "wb") as f:
            shutil.copyfileobj(file.file, f)
        try:
            new = _load_mapping(tmp)
        except Exception as e:
            raise HTTPException(400, detail=f"Not a usable MRP master - nothing changed: {e}")
        if new.empty:
            raise HTTPException(400, detail="The MRP master has no Department x Display x MRP rows - nothing changed")
        try:
            old = _load_mapping() if _find_mapping_file() else None
        except Exception:
            old = None
        archive = os.path.join(MAPPING_DIR, "Archive")
        os.makedirs(archive, exist_ok=True)
        stamp = datetime.datetime.now().strftime("%Y-%m-%d %H%M%S")
        archived = []
        for f in glob.glob(os.path.join(MAPPING_DIR, "*.xls*")):
            dest = os.path.join(archive, f"{stamp} {os.path.basename(f)}")
            shutil.move(f, dest)
            archived.append(os.path.basename(dest))
        shutil.move(tmp, os.path.join(MAPPING_DIR, name))
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    t = _targets(new)
    return {"ok": True, "file": name, "archived": archived,
            "rows": len(new), "groups": int(new.groupby(["DEPARTMENT", "DISPLAY"]).ngroups),
            "listed": int((new["MRP_LISTED"] > 0).sum()), "discontinued": int((new["MRP_LISTED"] == 0).sum()),
            "no_listed_groups": int(t.groupby(["DEPARTMENT", "DISPLAY"])["MRP_LISTED"].max().eq(0).sum()),
            "changes": _mapping_changes(old, new)}


@router.get("/mapping/download")
def download_mapping():
    """The active MRP master exactly as it was uploaded - edit it and import it back as the next version."""
    path = _find_mapping_file()
    if not path:
        raise HTTPException(404, detail=f"No MRP master in: {MAPPING_DIR}")
    return FileResponse(path, media_type=XLSX, filename=os.path.basename(path))
