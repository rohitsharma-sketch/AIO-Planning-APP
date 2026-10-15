"""Manual parity gate - the Sales Plan's rules checked against a planner's manual workbook (user, 2026-10-15: "ensure the
rules adopted now should give me accurate results in whichever type of variable session i will give to you for processing
against the manual one made by any human"; "i want foolproof rules ... mrps can be variable ... unless i add a rule or
subtract a rule this should not be deviating on every iteration").

The manual = the colleague's workbook layout ("MAMJ'27 DEPARTMENT - AOP (NEW APPROACH) - FOR BUYERS INPUT"):
  COMPARISON    cluster x department rows; row 4 names the blocks ("ACTUAL SALES VAL", "... (BUYERS INPUT)"), row 5 the
                half-months ("MAR'26 P1" ...). Columns are found by that text, so any months / column moves work.
  AOP           every store x CAT: "<Mon'YY> P1 AOP" / "P2 AOP", Ref Name, REV TAG (SSG / OTHER)
  STORE MASTER  STORE_NAME, REFF_STORE_MONTH_MIX
A layout that doesn't match stops with the reason (no guessing).

  python tools/manual_parity.py <manual.xlsx> --import    load the manual's inputs the rules read: ly_halves.json (the
                                                          departments' LY P1 / P2 per cluster) and aop_halves.json (store
                                                          and store x division P1 / P2, month-mix references); backups kept
  python tools/manual_parity.py <manual.xlsx> [--xlsx out.xlsx]
                                                          compare the saved plan with the manual; exit 1 if a check fails
The checks (Rs lakh; a cell passes when it shows 0.00 at 2 decimals):
  P1  SSG cluster x department x month x half = the manual's buyers'-input plan (COMPARISON)
  P2  every store x month (all divisions) = the manual's AOP sheet
  P3  every non-SSG store x division x month x half = the manual's AOP sheet
  P4  the SSG grand total and each month = the manual
SSG store x division is not a check: the plan caps each SSG cluster x month over all divisions and lets the division
split follow the buyer's growth - as the manual's own working does (its AOP sheet is the target)."""
import argparse, collections, datetime, itertools, json, os, re, shutil, sys

import openpyxl
import pandas as pd

BE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BE)
TOL = 0.005
HALF = re.compile(r"^\s*([A-Z]{3})'(\d{2})\s+(P[12])\s*$", re.I)
AOPH = re.compile(r"^\s*([A-Z]{3})'(\d{2})\s+(P[12])\s+AOP\s*$", re.I)
APPAREL = ("KIDS", "LADIES", "MENS")


class LayoutError(Exception):
    pass


def _mon(m3, yy):
    return f"{m3[:1].upper()}{m3[1:3].lower()}'{yy}"


def _num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def read_manual(path: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sh in ("COMPARISON", "AOP", "STORE MASTER"):
        if sh not in wb.sheetnames:
            raise LayoutError(f"sheet '{sh}' is missing")
    head = list(itertools.islice(wb["COMPARISON"].iter_rows(values_only=True), 5))
    if len(head) < 5:
        raise LayoutError("COMPARISON has fewer than 5 header rows")
    r4, r5 = list(head[3]), list(head[4])
    hdr = {str(v).strip().upper(): j for j, v in enumerate(r5) if v}
    for k in ("CLUSTER", "DIVISION", "SECTION", "DEPARTMENT"):
        if k not in hdr:
            raise LayoutError(f"COMPARISON row 5 has no '{k}' header")
    starts = sorted((j, str(v).strip().upper()) for j, v in enumerate(r4) if v)

    def block(*words):   # the block whose row-4 name has every word (the final plan: "CAPPED AT CLUSTER" + "(BUYERS INPUT)")
        s = next(((j, t) for j, t in starts if all(w in t for w in words)), None)
        if not s:
            raise LayoutError(f"COMPARISON row 4 has no block named with {' + '.join(words)}")
        end = next((j for j, _ in starts if j > s[0]), len(r5))
        out = {}
        for j in range(s[0], end):
            m = HALF.match(str(r5[j] or ""))
            if m:
                out[(_mon(m.group(1), m.group(2)), m.group(3).upper())] = j
        if not out:
            raise LayoutError(f"block '{s[1]}' has no MON'YY P1 / P2 headers")
        return out
    ly_cols, plan_cols = block("ACTUAL SALES"), block("CAPPED AT CLUSTER", "(BUYERS INPUT)")
    attr_j = hdr.get("ATTRIBUTE1", hdr["SECTION"])
    cmp_rows = []
    for r in wb["COMPARISON"].iter_rows(min_row=6, values_only=True):
        if not r or not r[hdr["CLUSTER"]] or not r[hdr["DEPARTMENT"]]:
            continue
        cmp_rows.append({"cluster": str(r[hdr["CLUSTER"]]).strip(), "division": str(r[hdr["DIVISION"]]).strip().upper(),
                         "section": str(r[hdr["SECTION"]] or "").strip(), "dept": str(r[hdr["DEPARTMENT"]]).strip(),
                         "attr": str(r[attr_j] or "").strip(),
                         "ly": {k: _num(r[j]) for k, j in ly_cols.items()}, "plan": {k: _num(r[j]) for k, j in plan_cols.items()}})
    # AOP sheet: the header row starts with "Store Name"; the first copy of each "Mon'YY P1 AOP" column counts
    aop, ah = [], None
    for r in wb["AOP"].iter_rows(values_only=True):
        if ah is None:
            if r and str(r[0] or "").strip().upper() == "STORE NAME":
                ah = {}
                for j, v in enumerate(r):
                    t = str(v or "").strip()
                    m = AOPH.match(t)
                    if m:
                        ah.setdefault((_mon(m.group(1), m.group(2)), m.group(3).upper()), j)
                    elif t.upper() in ("STORE NAME", "REF NAME", "CLUSTER", "STORE TAG", "REV TAG", "CAT"):
                        ah.setdefault(t.upper(), j)
                if not any(isinstance(k, tuple) for k in ah) or "CAT" not in ah or "REV TAG" not in ah:
                    raise LayoutError("the AOP sheet's header has no 'Mon'YY P1 AOP' columns, CAT or REV TAG")
            continue
        if not r or not r[0]:
            continue
        g = lambda k: str(r[ah[k]] or "").strip() if k in ah else ""   # noqa: E731
        aop.append({"store": str(r[0]).strip().upper(), "ref": g("REF NAME").upper(), "cluster": g("CLUSTER"), "tag": g("STORE TAG"),
                    "group": g("REV TAG").upper(), "division": g("CAT").upper(),
                    "half": {k: _num(r[j]) for k, j in ah.items() if isinstance(k, tuple)}})
    if ah is None:
        raise LayoutError("the AOP sheet has no 'Store Name' header row")
    sm = list(wb["STORE MASTER"].iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(sm[:5]) if r and "STORE_NAME" in [str(x or "").strip() for x in r]), None)
    if hi is None:
        raise LayoutError("STORE MASTER has no STORE_NAME header")
    ix = {str(h).strip(): j for j, h in enumerate(sm[hi]) if h}
    month_ref = {}
    if "REFF_STORE_MONTH_MIX" in ix:
        for r in sm[hi + 1:]:
            s, m = str(r[ix["STORE_NAME"]] or "").strip().upper(), str(r[ix["REFF_STORE_MONTH_MIX"]] or "").strip().upper()
            if s and m and m not in ("-", "NONE") and m != s:
                month_ref[s] = m
    for a in aop:   # the AOP sheet's own Ref Name is what its P1 / P2 come from (NLD: master FGN, sheet BGI)
        if a["group"] != "SSG" and a["ref"] and a["ref"] != a["store"]:
            month_ref[a["store"]] = a["ref"]
    return {"cmp": cmp_rows, "aop": aop, "month_ref": month_ref, "source": os.path.basename(path)}


def build_inputs(man: dict) -> tuple:
    """(ly_halves, aop_halves) - the manual's half-month inputs in the engine's file formats."""
    from engines.art_nm_engine import DEPT_SPLITS
    shares, parents = {}, {}
    for r in man["cmp"]:
        cl, d = r["cluster"].upper(), r["dept"].upper()
        mm = {}
        for (m, h), v in r["ly"].items():
            if h == "P1":
                p2 = r["ly"].get((m, "P2"), 0.0)
                if v + p2 > 0:   # 0 and 1 kept: a month whose LY sits in one half (15 Oct fix)
                    mm[m] = v / (v + p2)
        if not mm:
            continue
        if d in DEPT_SPLITS:
            parents[(cl, d)] = mm
        else:
            shares.setdefault(cl, {})[d] = mm
    for (cl, d), mm in parents.items():   # pre-split names go to their parts (the BIS rule: the parent's split wins)
        for k in DEPT_SPLITS[d]:
            shares.setdefault(cl, {})[k] = mm
    st, dv = collections.defaultdict(lambda: [0.0, 0.0]), collections.defaultdict(lambda: [0.0, 0.0])
    for a in man["aop"]:
        if a["division"] not in APPAREL:
            continue
        for (m, h), v in a["half"].items():
            i = 0 if h == "P1" else 1
            st[(a["store"], m)][i] += v
            dv[(a["store"], a["division"], m)][i] += v
    s_sh, d_sh = {}, {}
    for (s, m), (p1, p2) in st.items():
        if p1 + p2 > 0:
            s_sh.setdefault(s, {})[m] = p1 / (p1 + p2)
    for (s, d, m), (p1, p2) in dv.items():
        if p1 + p2 > 0:
            d_sh.setdefault(s, {}).setdefault(d, {})[m] = p1 / (p1 + p2)
    now = datetime.datetime.now().isoformat(timespec="seconds")
    half = "P1 = days 1-15, P2 = 16th to month end"
    return ({"source": f"{man['source']}, COMPARISON actual sales by half-month", "imported_at": now, "halves": half, "shares": shares},
            {"source": f"{man['source']}, AOP sheet (P1 / P2 AOP)", "imported_at": now, "halves": half,
             "shares": s_sh, "div_shares": d_sh, "month_ref": man["month_ref"]})


def do_import(man: dict) -> None:
    import engines.dept_sales_engine as DS
    ly, ao = build_inputs(man)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    for path, data in ((DS.LY_HALVES_PATH, ly), (DS.AOP_HALVES_PATH, ao)):
        if os.path.exists(path):
            bk = os.path.join(BE, "backups", f"{os.path.basename(path)}.before-parity-import-{stamp}")
            os.makedirs(os.path.dirname(bk), exist_ok=True)
            shutil.copy2(path, bk)
            print("backup:", bk)
        with open(path + ".tmp", "w", encoding="utf-8") as fh:
            json.dump(data, fh, indent=1)
        os.replace(path + ".tmp", path)
    print(f"imported: {sum(len(v) for v in ly['shares'].values())} cluster x department LY halves, "
          f"{len(ao['shares'])} stores, {sum(len(v) for v in ao['div_shares'].values())} store x division, "
          f"{len(ao['month_ref'])} month-mix references")
    print("next: the plan regenerates on its own (8010 sees the new files), then Display Type + Qty, then this tool without --import")


def compare(man: dict, plan: dict) -> tuple:
    """(checks [{id, what, cells, off, largest, ok}], {name: DataFrame})"""
    from engines.art_nm_engine import DEPT_SPLITS
    from engines.dept_sales_engine import canon_store
    parent = {c: p for p, cs in DEPT_SPLITS.items() for c in cs}
    info = {}
    for r in man["cmp"]:
        info.setdefault(parent.get(r["dept"], r["dept"]), (r["division"], r["section"], r["attr"]))
    planned = set((plan.get("meta") or {}).get("ty_months_active") or [])
    both = lambda m, h: (m, "P1") in man["cmp"][0]["plan"] and (m, "P2") in man["cmp"][0]["plan"]   # noqa: E731
    months = sorted({m for (m, h) in man["cmp"][0]["plan"] if m in planned and both(m, h)}) if man["cmp"] else []
    skipped = sorted({m for (m, h) in (man["cmp"][0]["plan"] if man["cmp"] else {}) if m not in months})
    if skipped:
        print("not compared (the manual has them, the plan doesn't plan them or only one half):", ", ".join(skipped))
    a = collections.Counter()
    for r in man["cmp"]:
        for (m, h), v in r["plan"].items():
            if m in months:
                a[("Manual", r["cluster"], parent.get(r["dept"], r["dept"]), m, h)] += v
    for s, sd in plan["stores"].items():
        if not sd.get("is_ssg"):
            continue
        for dd in sd["divisions"].values():
            for m, md in dd["months"].items():
                if m in months:
                    for d, x in md["departments"].items():
                        if x.get("active"):
                            k = (sd.get("cluster") or "", parent.get(d, d), m)
                            a[("Portal",) + k + ("P1",)] += x.get("ty_p1", 0.0)
                            a[("Portal",) + k + ("P2",)] += x.get("ty_p2", 0.0)
    t1 = pd.DataFrame([(*k, v) for k, v in a.items()], columns=["src", "Cluster", "Department", "Month", "Fortnight", "v"]
                      ).pivot_table(index=["Cluster", "Department", "Month", "Fortnight"], columns="src", values="v",
                                    aggfunc="sum", fill_value=0.0).reset_index()
    b, lab = collections.Counter(), {}
    for r in man["aop"]:
        s = canon_store(r["store"])
        lab[s] = (r["cluster"], r["tag"], "SSG" if r["group"] == "SSG" else "OTHER")
        for (m, h), v in r["half"].items():
            if m in months and r["division"] in APPAREL:
                b[("Manual", s, r["division"], m, h)] += v
    for s, sd in plan["stores"].items():
        lab.setdefault(s, (sd.get("cluster") or "", sd.get("tag") or "", "SSG" if sd.get("is_ssg") else "OTHER"))
        for dvn, dd in sd["divisions"].items():
            if dvn in APPAREL:
                for m, md in dd["months"].items():
                    if m in months:
                        act = [x for x in md["departments"].values() if x.get("active")]
                        b[("Portal", s, dvn, m, "P1")] += sum(x.get("ty_p1", 0.0) for x in act)
                        b[("Portal", s, dvn, m, "P2")] += sum(x.get("ty_p2", 0.0) for x in act)
    t2 = pd.DataFrame([(*k, v) for k, v in b.items()], columns=["src", "Store", "Division", "Month", "Fortnight", "v"]
                      ).pivot_table(index=["Store", "Division", "Month", "Fortnight"], columns="src", values="v",
                                    aggfunc="sum", fill_value=0.0).reset_index()
    for t in (t1, t2):
        for c in ("Manual", "Portal"):
            if c not in t:
                t[c] = 0.0
        t["Diff"] = t["Portal"] - t["Manual"]
    t1["Division"], t1["Sub-division"], t1["Attribute"] = zip(*[info.get(d, ("", "", "")) for d in t1["Department"]])
    t2["Cluster"], t2["Store tag"], t2["Group"] = zip(*[lab.get(s, ("", "", "")) for s in t2["Store"]])

    def chk(cid, what, g):
        d = (g["Portal"] - g["Manual"]).abs()
        return {"id": cid, "what": what, "cells": int(len(d)), "off": int((d >= TOL).sum()),
                "largest": float(d.max()) if len(d) else 0.0, "ok": bool((d < TOL).all())}
    tot = t1.groupby("Month", as_index=False)[["Manual", "Portal"]].sum()
    tot = pd.concat([tot, pd.DataFrame([{"Month": "Total", "Manual": tot.Manual.sum(), "Portal": tot.Portal.sum()}])])
    checks = [chk("P1", "SSG cluster x department x month x half = the manual's buyers'-input plan", t1),
              chk("P2", "Every store x month (all divisions) = the manual's AOP sheet",
                  t2.groupby(["Store", "Month"], as_index=False)[["Manual", "Portal"]].sum()),
              chk("P3", "Every non-SSG store x division x month x half = the manual's AOP sheet", t2[t2.Group != "SSG"]),
              chk("P4", "SSG grand total and each month = the manual", tot)]
    return checks, {"SSG cluster x dept x half": t1, "All stores x div x half": t2}


def write_xlsx(out: str, checks: list, frames: dict) -> None:
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        pd.DataFrame([{"Check": c["id"], "What": c["what"], "Result": "PASS" if c["ok"] else "FAIL", "Cells": c["cells"],
                       "Not 0.00": c["off"], "Largest diff (Rs L)": round(c["largest"], 4)} for c in checks]
                     ).to_excel(w, sheet_name="Checks", index=False)
        for nm, t in frames.items():
            dims = [c for c in ("Group", "Store tag", "Cluster", "Store", "Division", "Sub-division", "Attribute", "Department",
                                "Month", "Fortnight") if c in t]
            summ = []
            for k in range(len(dims) + 1):
                for keys in itertools.combinations(dims, k):
                    g = t.groupby(list(keys))[["Manual", "Portal"]].sum() if keys else t[["Manual", "Portal"]].sum().to_frame().T
                    d = (g["Portal"] - g["Manual"]).abs()
                    summ.append({"Levels": " x ".join(keys) or "Grand total", "Rows": len(g),
                                 "Max |Diff| (Rs L)": float(d.max()), "Rows not 0.00": int((d >= TOL).sum())})
            pd.DataFrame(summ).to_excel(w, sheet_name=(nm + " levels")[:31], index=False)
            t[dims + ["Manual", "Portal", "Diff"]].to_excel(w, sheet_name=nm[:31], index=False)
    print("written", out)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Sales Plan vs a planner's manual workbook")
    ap.add_argument("manual")
    ap.add_argument("--import", dest="imp", action="store_true", help="load the manual's half-month inputs first")
    ap.add_argument("--xlsx", help="also write the comparison workbook here")
    a = ap.parse_args(argv)
    try:
        man = read_manual(a.manual)
    except LayoutError as e:
        print("THE MANUAL'S LAYOUT DOESN'T MATCH:", e)
        return 2
    if a.imp:
        do_import(man)
        return 0
    import engines.dept_sales_engine as DS
    from plan_cache import load_json
    checks, frames = compare(man, load_json(DS.FINAL_PLAN_PATH))
    for c in checks:
        print(f"{'PASS' if c['ok'] else 'FAIL'}  {c['id']}  {c['what']}: {c['off']} of {c['cells']:,} not 0.00 "
              f"(largest {c['largest']:.4f} L)")
    if a.xlsx:
        write_xlsx(a.xlsx, checks, frames)
    return 0 if all(c["ok"] for c in checks) else 1


if __name__ == "__main__":
    sys.exit(main())
