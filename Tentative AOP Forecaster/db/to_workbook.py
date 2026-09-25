"""
Builds an inputs.xlsx from Postgres only — the DB-backed equivalent of
config_store.build_workbook(). engine_v3.py is completely unchanged: it still
reads an .xlsx with the same sheet layout it always has (skiprows=3 etc.);
only where that file's data comes from changes (Postgres, not levers.json).

Sheet layout notes (matching engine_v3.load_inputs()):
  - Store Master, Store Actuals, Growth %, NSO Opening Months: 3 blank
    preamble rows, header on row 4 (skiprows=3) — the preamble content is
    never parsed, so it's left blank here rather than copying decorative
    title text from the old workbook.
  - AOP (Optional): header on row 1, row 2 is a skipped note row.
"""
import datetime
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from db.models.masterdata import NsoOpening, Store
from db.models.planning_inputs import InputValue, Period
from engine_v3 import auto_tag

FY27_M = ["Mar'26", "Apr'26", "May'26", "Jun'26", "Jul'26", "Aug'26", "Sep'26",
          "Oct'26", "Nov'26", "Dec'26", "Jan'27", "Feb'27", "Mar'27"]
FY28_M = ["Mar'27", "Apr'27", "May'27", "Jun'27", "Jul'27", "Aug'27", "Sep'27",
          "Oct'27", "Nov'27", "Dec'27", "Jan'28", "Feb'28", "Mar'28"]
DIVS = ["GM", "KIDS", "LADIES", "MENS", "RETAIL"]
MON_NAMES = {1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "May", 6: "Jun",
             7: "Jul", 8: "Aug", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dec"}


def _date_to_label(d: datetime.date) -> str:
    return f"{MON_NAMES[d.month]}'{d.year % 100:02d}"


def _write_sheet(wb, name, columns, rows, preamble_rows=3, post_note=None):
    ws = wb.create_sheet(name[:31])
    r = 1
    for _ in range(preamble_rows):
        r += 1
    for c, h in enumerate(columns, 1):
        ws.cell(r, c, h)
    r += 1
    if post_note is not None:
        ws.cell(r, 1, post_note)
        r += 1
    for row in rows:
        for c, v in enumerate(row, 1):
            if v is not None:
                ws.cell(r, c, v)
        r += 1


def build_workbook_from_db(session: Session, out_path: str) -> str:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Store Master — current rows only (valid_to IS NULL), and only stores that
    # are actually in the planning universe (i.e. carry a Tag from Store
    # Master.xlsx) — masterdata.stores also holds ERP-only placeholder/inactive
    # codes from site_master (e.g. "BIHAR-DUMMY") that were never part of it.
    stores = session.execute(
        select(Store).where(Store.valid_to.is_(None), Store.tag.is_not(None))
    ).scalars().all()
    _write_sheet(wb, "Store Master", ["Store", "Ref Store", "Cluster", "Tag"],
                 [[s.store_id, s.ref_store, s.cluster_key, auto_tag(s.tag, s.opening_date, s.store_current_status)]
                  for s in stores])

    # Store Actuals — pivot input_values(store_actuals) to Store x Division x
    # Attribute x FY27 month grid. row_key carries the ATTRIBUTE1 value the sync
    # wrote (see store_actuals_sync.py) - grouping includes it now so two rows
    # for the same store/division/month but different attributes don't
    # collapse into one one (the pre-attribute pivot only grouped by
    # (store, division), which silently overwrote one attribute's value with
    # another's once store_actuals_sync started writing more than one row_key
    # per store/division/month). A blank row_key (pre-migration data, or a
    # lever other than store_actuals ever reusing this path) still pivots
    # fine - it's just one more attribute value, "".
    periods = {p.period_id: p.label for p in session.execute(select(Period)).scalars().all()}
    actuals = session.execute(select(InputValue).where(InputValue.lever_key == "store_actuals")).scalars().all()
    grid = {}
    for v in actuals:
        grid.setdefault((v.store_id, v.division_code, v.row_key), {})[periods.get(v.period_id)] = float(v.value)
    actuals_rows = [[store, div, attr] + [vals.get(m) for m in FY27_M] for (store, div, attr), vals in grid.items()]
    _write_sheet(wb, "Store Actuals", ["Store", "Division", "Attribute"] + FY27_M, actuals_rows)

    # Store Actuals FY26 — raw FY26 month actuals (store_actuals_sync step 5),
    # the engine's placeholder base for a not-yet-closed FY27 month.
    fy26_m = [f"{m[:4]}{int(m[4:]) - 1:02d}" for m in FY27_M[:-1]]  # Mar'25..Feb'26
    fy26_grid = {}
    for v in session.execute(select(InputValue).where(InputValue.lever_key == "store_actuals_fy26")).scalars().all():
        fy26_grid.setdefault((v.store_id, v.division_code), {})[periods.get(v.period_id)] = float(v.value)
    _write_sheet(wb, "Store Actuals FY26", ["Store", "Division"] + fy26_m,
                 [[store, div] + [vals.get(m) for m in fy26_m] for (store, div), vals in fy26_grid.items()])

    # Growth % — pivot row_key (OVERALL/division) x FY28 month
    growth = session.execute(select(InputValue).where(InputValue.lever_key == "growth_pct")).scalars().all()
    ggrid = {}
    for v in growth:
        ggrid.setdefault(v.row_key, {})[periods.get(v.period_id)] = float(v.value)
    growth_rows = [[key] + [vals.get(m) for m in FY28_M[1:]] for key, vals in ggrid.items()]  # FY28_M[1:] = Apr'27..Mar'28
    _write_sheet(wb, "Growth %", ["Division"] + FY28_M[1:], growth_rows)

    # NSO Opening Months — unnamed only (is_named=False)
    nso = session.execute(select(NsoOpening).where(NsoOpening.is_named.is_(False))).scalars().all()
    _write_sheet(wb, "NSO Opening Months", ["Store Code", "Opening Month"],
                 [[n.store_id, _date_to_label(n.opening_month)] for n in nso])

    # AOP (Optional) — store x division x FY28 month (header on row 1, note on row 2)
    aop = session.execute(select(InputValue).where(InputValue.lever_key == "aop_optional")).scalars().all()
    agrid = {}
    for v in aop:
        agrid.setdefault((v.store_id, v.division_code), {})[periods.get(v.period_id)] = float(v.value)
    aop_rows = [[store, div] + [vals.get(m) for m in FY28_M] for (store, div), vals in agrid.items()]
    _write_sheet(wb, "AOP (Optional)", ["Store", "Division"] + FY28_M, aop_rows,
                 preamble_rows=0, post_note="Migrated from Postgres — see db/to_workbook.py")

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path
