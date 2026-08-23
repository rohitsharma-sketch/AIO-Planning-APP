"""
Configuration layer for the AOP Forecaster.

Instead of treating inputs.xlsx as a single monolithic template that must be
re-uploaded every run, each sheet ("lever") is stored as its own editable
element list in config/levers.json.  The workbook layout of every lever
(title / note rows, header row position, columns) is preserved exactly so the
engine's fixed parsing (skiprows=3, Mar27 Targets header=0 / skip row 2) keeps
working when a workbook is rebuilt from the configuration.

Public API
----------
load_config()                -> dict            (auto-seeds from inputs.xlsx if empty)
import_workbook(path, name)  -> dict            (xlsx → becomes the default for ALL levers)
update_lever(key, rows)      -> lever dict      (editor save: replace the element list)
reset_config()               -> dict            (re-seed from the project inputs.xlsx)
build_workbook(out_path)     -> out_path        (config → inputs.xlsx for the engine)
"""
import datetime
import json
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

_HERE       = os.path.dirname(os.path.abspath(__file__))
CONFIG_DIR  = os.path.join(_HERE, "config")
CONFIG_FILE = os.path.join(CONFIG_DIR, "levers.json")
SEED_FILE   = os.path.join(_HERE, "inputs.xlsx")

# Levers the engine actually reads; everything else in the workbook is carried along as optional.
REQUIRED_LEVERS = {
    "Store Master":       "One row per store: Store, Ref Store, Cluster, Tag",
    "Store Actuals":      "Store × Division FY27 day-shifted sales (₹ Lakhs) — synced from the Calendar Engine",
    "Growth %":           "Frozen at a flat 6% — adjusted later in the Review step",
    "NSO Opening Months": "Unnamed NSO stores and their go-live month",
}
OPTIONAL_LEVERS = {
    "AOP (Optional)":       "Per store×division AOP overrides for any FY28 month, plus the Mar'27 actuals fallback",
    "Named NSO (Optional)": "Opening-month overrides for named NSO stores",
    "Settings":             "Forecaster settings (paths) — informational",
}


# Growth % is frozen at a flat rate in the configuration; it is re-adjusted later in the Review step.
FROZEN_GROWTH_PCT = 6.0
GROWTH_KEY = "growth_pct"
GROWTH_DIVS = ["GM", "KIDS", "LADIES", "MENS", "RETAIL"]


def freeze_growth(cfg, pct=FROZEN_GROWTH_PCT):
    """OVERALL = pct for every month, division rows blank (inherit OVERALL). Note rows untouched."""
    for l in cfg.get("levers", []):
        if l["key"] != GROWTH_KEY:
            continue
        ncol = len(l["columns"])
        names = {str(r[0]).strip().upper() for r in l["rows"] if r and r[0] is not None}
        if "OVERALL" not in names:
            l["rows"].insert(0, ["OVERALL"] + [None] * (ncol - 1))
        for r in l["rows"]:
            head = str(r[0]).strip().upper() if r and r[0] is not None else ""
            if head == "OVERALL":
                for i in range(1, ncol): r[i] = pct
            elif head in GROWTH_DIVS:
                for i in range(1, ncol): r[i] = None
        l["frozen"] = pct
    return cfg


def _key(sheet):
    return sheet.lower().replace("%", "pct").replace("(", "").replace(")", "").strip().replace(" ", "_")


def _cell(v):
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


# ── Parsing ──────────────────────────────────────────────────────────────────
def _parse_sheet(ws):
    """Detect the header row (densest of the first 4 rows) and split the sheet into
    preamble rows / header / optional post-header note / data rows."""
    rows = [[_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
    # trim fully-empty trailing columns
    width = 0
    for r in rows:
        for i, v in enumerate(r):
            if v is not None and str(v).strip() != "":
                width = max(width, i + 1)
    rows = [r[:width] for r in rows]
    if not rows:
        return None

    def density(r):
        return sum(1 for v in r if v is not None and str(v).strip() != "")

    probe = rows[:4]
    header_idx = max(range(len(probe)), key=lambda i: (density(probe[i]), -i))
    header = [str(v).strip() if v is not None else "" for v in rows[header_idx]]
    # drop trailing blank header cells
    while header and header[-1] == "":
        header.pop()
    ncol = len(header)
    preamble = [[(str(v) if v is not None else None) for v in r[:ncol]] + [None] * (ncol - len(r)) for r in rows[:header_idx]]

    body = rows[header_idx + 1:]
    post_note = None
    if body and density(body[0]) == 1 and isinstance(body[0][0], str) and len(body[0][0]) > 20 and density(body[0]) < max(2, ncol // 2):
        post_note = body[0][0]
        body = body[1:]

    data = []
    for r in body:
        r = list(r[:ncol]) + [None] * (ncol - len(r))
        if density(r) == 0:
            continue
        data.append(r)
    return {
        "preamble":  preamble,     # rows above the header (titles / notes / blanks), positions preserved
        "header_row": header_idx + 1,
        "columns":   header,
        "post_note": post_note,    # single note line directly under the header (Mar27 Targets style)
        "rows":      data,
    }


def _drop_column(lever, name):
    """Remove a column (and its cells) from a lever if present."""
    norm = lambda x: str(x).strip().lower()
    cols = [norm(c) for c in lever["columns"]]
    if norm(name) not in cols:
        return
    i = cols.index(norm(name))
    lever["columns"].pop(i)
    lever["rows"] = [r[:i] + r[i + 1:] for r in lever["rows"]]
    lever["preamble"] = [p[:i] + p[i + 1:] if len(p) > i else p for p in lever.get("preamble", [])]


_AOP_DIVS  = ["GM", "KIDS", "LADIES", "MENS", "RETAIL"]
_AOP_MONTHS = ["Mar'27","Apr'27","May'27","Jun'27","Jul'27","Aug'27","Sep'27",
               "Oct'27","Nov'27","Dec'27","Jan'28","Feb'28","Mar'28"]
_AOP_NOTE = ("Optional. Mar'27 column = fallback actual when the real Mar'27 sales is missing/zero "
             "(feeds the Mar'28 base). Apr'27..Mar'28 columns = direct AOP overrides — type any "
             "value to force that store×division×month's final forecast; leave blank everywhere "
             "else to let the engine compute it. Values in Rs Lakhs.")


def _migrate_mar27_wide_to_aop(parsed):
    """Legacy 'Mar27 Targets' sheet (Store + one column per division, Mar'27 only) →
    the long-format 'AOP (Optional)' shape (Store, Division, every FY28 month) so the
    same lever also accepts overrides for any other month. Only nonzero legacy values
    are carried over, written into the Mar'27 column; other months start blank."""
    cols = parsed["columns"]
    div_idx = {d: cols.index(d) for d in _AOP_DIVS if d in cols}
    if "Store" not in cols or not div_idx:
        return parsed   # unrecognised shape — leave as-is
    si = cols.index("Store")
    new_rows = []
    for r in parsed["rows"]:
        store = r[si] if si < len(r) else None
        if store is None:
            continue
        for d, ci in div_idx.items():
            v = r[ci] if ci < len(r) else None
            if v not in (None, 0):
                row = [None] * len(_AOP_MONTHS)
                row[0] = v   # Mar'27 is column 0
                new_rows.append([store, d] + row)
    return {
        "preamble": [], "header_row": 1, "columns": ["Store", "Division"] + _AOP_MONTHS,
        "post_note": _AOP_NOTE, "rows": new_rows,
    }


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    levers = []
    for ws in wb.worksheets:
        parsed = _parse_sheet(ws)
        if not parsed:
            continue
        sheet = ws.title
        if sheet == "Store Master":
            _drop_column(parsed, "CONC")   # legacy column — engine derives its own keys
        if sheet == "Mar27 Targets":
            # Renamed lever: same data, reshaped so any FY28 month can hold an AOP override.
            sheet = "AOP (Optional)"
            parsed = _migrate_mar27_wide_to_aop(parsed)
        levers.append({
            "key":         _key(sheet),
            "sheet":       sheet,
            "required":    sheet in REQUIRED_LEVERS,
            "description": REQUIRED_LEVERS.get(sheet) or OPTIONAL_LEVERS.get(sheet) or "Carried through to the workbook unchanged",
            **parsed,
        })
    wb.close()
    missing = [s for s in REQUIRED_LEVERS if s not in {l["sheet"] for l in levers}]
    if missing:
        raise ValueError(f"Workbook is missing required sheet(s): {', '.join(missing)}")
    return levers


# ── Persistence ──────────────────────────────────────────────────────────────
def _now():
    return datetime.datetime.now().isoformat(timespec="seconds")


def _save(cfg):
    os.makedirs(CONFIG_DIR, exist_ok=True)
    tmp = CONFIG_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(cfg, f)
    os.replace(tmp, CONFIG_FILE)
    return cfg


def load_config():
    if os.path.exists(CONFIG_FILE):
        with open(CONFIG_FILE, encoding="utf-8") as f:
            return json.load(f)
    if os.path.exists(SEED_FILE):
        return import_workbook(SEED_FILE, source="inputs.xlsx (project default)")
    return _save({"levers": [], "source": None, "updated": _now()})


def import_workbook(path, source=None):
    levers = parse_workbook(path)
    stamp = _now()
    for l in levers:
        l["updated"] = stamp
    return _save(freeze_growth({"levers": levers, "source": source or os.path.basename(path), "imported": stamp, "updated": stamp}))


def reset_config():
    if not os.path.exists(SEED_FILE):
        raise FileNotFoundError("Project inputs.xlsx not found — nothing to reset to")
    return import_workbook(SEED_FILE, source="inputs.xlsx (project default)")


def update_lever(key, rows=None, columns=None):
    cfg = load_config()
    for l in cfg["levers"]:
        if l["key"] == key:
            if columns is not None:
                l["columns"] = [str(c).strip() for c in columns]
            ncol = len(l["columns"])
            if rows is not None:
                clean = []
                for r in rows:
                    r = list(r)[:ncol] + [None] * (ncol - len(r))
                    r = [None if (v == "" or v is None) else v for v in r]
                    if all(v is None for v in r):
                        continue
                    clean.append(r)
                l["rows"] = clean
            l["updated"] = _now()
            cfg["updated"] = l["updated"]
            _save(cfg)
            return l
    raise KeyError(key)


def summary(cfg):
    """Lightweight view for the upload page."""
    return {
        "source":   cfg.get("source"),
        "imported": cfg.get("imported"),
        "updated":  cfg.get("updated"),
        "calendar_sync": cfg.get("calendar_sync") or {},
        "levers": [{
            "key": l["key"], "sheet": l["sheet"], "required": l["required"], "description": l["description"],
            "columns": l["columns"], "n_rows": len(l["rows"]), "updated": l.get("updated"),
            "frozen": l.get("frozen"),
        } for l in cfg["levers"]],
        "frozen_growth_pct": FROZEN_GROWTH_PCT,
    }


# ── Workbook generation ──────────────────────────────────────────────────────
def build_workbook(out_path, cfg=None):
    cfg = freeze_growth(cfg or load_config())     # sessions always start from the frozen growth rate
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    h_fill = PatternFill("solid", fgColor="1a2744")
    h_font = Font(bold=True, color="FFFFFF")
    t_font = Font(bold=True, size=12)
    for l in cfg["levers"]:
        ws = wb.create_sheet(l["sheet"][:31])
        r = 1
        for pre in l["preamble"]:
            for c, v in enumerate(pre, 1):
                if v is not None:
                    cell = ws.cell(r, c, v)
                    if r == 1:
                        cell.font = t_font
            r += 1
        assert r == l["header_row"], f"{l['sheet']}: header row drifted"
        for c, h in enumerate(l["columns"], 1):
            cell = ws.cell(r, c, h)
            cell.fill = h_fill; cell.font = h_font; cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = max(len(str(h)) + 4, 12)
        r += 1
        if l.get("post_note"):
            ws.cell(r, 1, l["post_note"]).font = Font(italic=True, color="666666")
            r += 1
        for row in l["rows"]:
            for c, v in enumerate(row, 1):
                if v is not None:
                    ws.cell(r, c, _coerce(v))
            r += 1
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path


def _coerce(v):
    """Editor sends strings; turn numeric-looking strings back into numbers."""
    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return None
        try:
            if s.lstrip("-").replace(".", "", 1).isdigit():
                return float(s) if "." in s else int(s)
        except ValueError:
            pass
        return s
    return v


# ── Unified store configuration (Store Master + NSO Opening Months + Named NSO) ──
STORE_LEVERS = {"store_master": "Store Master", "nso_opening_months": "NSO Opening Months", "named_nso_optional": "Named NSO (Optional)"}
UNNAMED_PREFIXES = ("NS-", "AD-", "MAMJ")


def _col(lever, *names):
    """Index of the first column whose header matches one of `names` (case/space-insensitive)."""
    norm = lambda s: str(s).strip().lower().replace("_", " ")
    cols = [norm(c) for c in lever["columns"]]
    for n in names:
        if norm(n) in cols:
            return cols.index(norm(n))
    return None


def _lever(cfg, key):
    return next((l for l in cfg["levers"] if l["key"] == key), None)


def _get(row, i):
    if i is None or i >= len(row):
        return None
    v = row[i]
    return None if v is None or str(v).strip() == "" else (v.strip() if isinstance(v, str) else v)


def get_stores(cfg=None):
    """Merge the three levers into one record per store."""
    cfg = cfg or load_config()
    sm, nso, named = (_lever(cfg, k) for k in ("store_master", "nso_opening_months", "named_nso_optional"))
    stores, order = {}, []

    def rec(code):
        if code not in stores:
            stores[code] = {"store": code, "ref_store": None, "cluster": None, "tag": None,
                            "opening_month": None, "notes": None, "in_master": False, "nso_source": None,
                            "order_master": None, "order_nso": None, "order_named": None}
            order.append(code)
        return stores[code]

    if sm:
        iS, iR, iCl, iT = (_col(sm, "Store"), _col(sm, "Ref Store"), _col(sm, "Cluster"), _col(sm, "Tag", "Store Tag"))
        for n, r in enumerate(sm["rows"]):
            code = _get(r, iS)
            if not code: continue
            s = rec(str(code)); s.update(in_master=True, order_master=n, ref_store=_get(r, iR), cluster=_get(r, iCl), tag=_get(r, iT))
    if nso:
        iS, iO, iT, iR, iN = (_col(nso, "Store Code", "Store"), _col(nso, "Opening Month"), _col(nso, "Store Tag", "Tag"), _col(nso, "Ref Store"), _col(nso, "Notes"))
        for n, r in enumerate(nso["rows"]):
            code = _get(r, iS)
            if not code: continue
            s = rec(str(code)); s["nso_source"] = "unnamed"; s["order_nso"] = n; s["opening_month"] = _get(r, iO); s["notes"] = _get(r, iN)
            s["tag"] = s["tag"] or _get(r, iT); s["ref_store"] = s["ref_store"] or _get(r, iR)
    if named:
        iS, iR, iT, iO, iCl, iN = (_col(named, "Store"), _col(named, "Ref Store"), _col(named, "Store Tag", "Tag"), _col(named, "Opening Month"), _col(named, "Cluster"), _col(named, "Notes"))
        for n, r in enumerate(named["rows"]):
            code = _get(r, iS)
            if not code: continue
            s = rec(str(code)); s["order_named"] = n
            if s["nso_source"] is None: s["nso_source"] = "named"
            s["opening_month"] = s["opening_month"] or _get(r, iO); s["notes"] = s["notes"] or _get(r, iN)
            s["tag"] = s["tag"] or _get(r, iT); s["ref_store"] = s["ref_store"] or _get(r, iR); s["cluster"] = s["cluster"] or _get(r, iCl)

    tags = sorted({s["tag"] for s in stores.values() if s["tag"]})
    clusters = sorted({s["cluster"] for s in stores.values() if s["cluster"]})
    settings = _lever(cfg, "settings")
    return {
        "stores": [stores[c] for c in order],
        "tags": tags, "clusters": clusters,
        "months": ["Mar'27","Apr'27","May'27","Jun'27","Jul'27","Aug'27","Sep'27","Oct'27","Nov'27","Dec'27","Jan'28","Feb'28","Mar'28"],
        "settings": {"columns": settings["columns"], "rows": settings["rows"]} if settings else None,
        "updated": cfg.get("updated"),
    }


def put_stores(stores, settings_rows=None):
    """Split unified store records back into the three levers (and optionally Settings)."""
    cfg = load_config()
    sm, nso, named = (_lever(cfg, k) for k in ("store_master", "nso_opening_months", "named_nso_optional"))
    stamp = _now()

    def row_for(lever, mapping):
        r = [None] * len(lever["columns"])
        for names, val in mapping:
            i = _col(lever, *names)
            if i is not None: r[i] = val
        return r

    def is_unnamed(code):
        return str(code).upper().startswith(UNNAMED_PREFIXES)

    clean = []
    for s in stores:
        code = (s.get("store") or "").strip()
        if not code: continue
        s = {**s, "store": code}
        for k in ("ref_store", "cluster", "tag", "opening_month", "notes"):
            v = s.get(k); s[k] = v.strip() if isinstance(v, str) and v.strip() else (v if not isinstance(v, str) else None)
        if s.get("in_master") is None: s["in_master"] = True
        if s.get("nso_source") is None and s.get("opening_month"):
            s["nso_source"] = "unnamed" if is_unnamed(code) else "named"
        clean.append(s)

    BIG = 10 ** 9
    def ordered(items, key):
        return sorted(items, key=lambda s: (s.get(key) if s.get(key) is not None else BIG, s["store"]))

    if sm:
        _drop_column(sm, "CONC")
        sm["rows"] = [row_for(sm, [(("Store",), s["store"]), (("Ref Store",), s["ref_store"]),
                                   (("Cluster",), s["cluster"]), (("Tag", "Store Tag"), s["tag"])]) for s in ordered([s for s in clean if s["in_master"]], "order_master")]
        sm["updated"] = stamp
    if nso:
        nso["rows"] = [row_for(nso, [(("Store Code", "Store"), s["store"]), (("Opening Month",), s["opening_month"]), (("Store Tag", "Tag"), s["tag"]),
                                     (("Ref Store",), s["ref_store"]), (("Notes",), s["notes"])]) for s in ordered([s for s in clean if s["nso_source"] == "unnamed"], "order_nso")]
        nso["updated"] = stamp
    if named:
        named["rows"] = [row_for(named, [(("Store",), s["store"]), (("Ref Store",), s["ref_store"]), (("Store Tag", "Tag"), s["tag"]),
                                         (("Opening Month",), s["opening_month"]), (("Cluster",), s["cluster"]), (("Notes",), s["notes"])]) for s in ordered([s for s in clean if s["nso_source"] == "named"], "order_named")]
        named["updated"] = stamp
    st = _lever(cfg, "settings")
    if st and settings_rows is not None:
        st["rows"] = [list(r)[:len(st["columns"])] + [None] * (len(st["columns"]) - len(r)) for r in settings_rows if any(v not in (None, "") for v in r)]
        st["updated"] = stamp
    cfg["updated"] = stamp
    _save(cfg)
    return get_stores(cfg)
